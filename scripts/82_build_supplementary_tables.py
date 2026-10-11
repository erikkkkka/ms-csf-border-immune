#!/usr/bin/env python3
"""Build indexed, submission-ready supplementary table workbooks."""

from __future__ import annotations

import argparse
import hashlib
import math
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
NAVY = "243746"
BLUE = "4E79A7"
PALE_BLUE = "EAF0F5"
PALE_GREY = "F3F5F6"
MID_GREY = "66717A"
WHITE = "FFFFFF"


@dataclass(frozen=True)
class SheetSpec:
    name: str
    source: str | tuple[str, ...]
    description: str
    transform: Callable[[pd.DataFrame], pd.DataFrame] | None = None


@dataclass(frozen=True)
class WorkbookSpec:
    filename: str
    title: str
    scope: str
    sheets: tuple[SheetSpec, ...]


def combine_with_dataset(root: Path, paths: tuple[str, ...]) -> pd.DataFrame:
    frames = []
    for relative in paths:
        frame = pd.read_csv(root / relative)
        if "dataset" not in frame.columns and "gse" not in frame.columns:
            frame.insert(0, "dataset", Path(relative).name.split("_")[0])
        frames.append(frame)
    return pd.concat(frames, ignore_index=True, sort=False)


def specs() -> tuple[WorkbookSpec, ...]:
    return (
        WorkbookSpec(
            "Supplementary_Table_1_Cohorts_and_QC.xlsx",
            "Supplementary Table 1 | Cohorts, sample metadata and quality control",
            "Cohort provenance, sample-level metadata, cell retention, ambient-RNA correction and independent doublet calls.",
            (
                SheetSpec("Cohort_manifest", "metadata/cohort_roles.csv", "Dataset accessions, analysis roles and external-cohort inclusion."),
                SheetSpec("Discovery_metadata", "metadata/all_geo_sample_metadata.csv", "Harmonized discovery-cohort sample metadata."),
                SheetSpec("External_metadata", "metadata/GSE328265_sample_metadata.csv", "Independent external cohort sample metadata."),
                SheetSpec("Sample_count_summary", "metadata/sample_count_summary.csv", "Sample and group counts used to audit metadata harmonization."),
                SheetSpec("QC_GSE138266", "results/qc/GSE138266_cleanup_sample_summary.csv", "Sample-level final QC, DecontX and doublet-union summary."),
                SheetSpec("QC_GSE194078", "results/qc/GSE194078_cleanup_sample_summary.csv", "Sample-level final QC, DecontX and doublet-union summary."),
                SheetSpec("QC_GSE286068", "results/qc/GSE286068_cleanup_sample_summary.csv", "Sample-level final QC, DecontX and doublet-union summary."),
                SheetSpec("QC_GSE328265", "results/qc/GSE328265_cleanup_sample_summary.csv", "Sample-level final QC, DecontX and doublet-union summary."),
            ),
        ),
        WorkbookSpec(
            "Supplementary_Table_2_Annotation_and_State.xlsx",
            "Supplementary Table 2 | Myeloid annotation and state abundance",
            "Revised CSF myeloid labels, BAM-like robustness audits and sample-level lineage-state abundance.",
            (
                SheetSpec("Revised_state_summary", "results/figure1c_annotation_audit/Figure1c_multicohort_revised_state_summary.csv", "Multicohort revised state names and marker evidence used in Figure 1c."),
                SheetSpec("Annotation_metrics", "results/figure1c_annotation_audit/Figure1c_multicohort_annotation_metrics.csv", "Silhouette and confidence-tier annotation audit."),
                SheetSpec("BAM_robustness", "results/bam_robustness/BAM_robustness_summary.csv", "Stable-core retention and annotation robustness summary."),
                SheetSpec("BAM_marker_expression", (
                    "results/bam_robustness/GSE138266_BAM_marker_expression.csv",
                    "results/bam_robustness/GSE194078_BAM_marker_expression.csv",
                    "results/bam_robustness/GSE286068_BAM_marker_expression.csv",
                ), "Cohort-specific expression of BAM-like annotation markers."),
                SheetSpec("Marker_leave_one_out", (
                    "results/bam_robustness/GSE138266_BAM_marker_leave_one_out.csv",
                    "results/bam_robustness/GSE194078_BAM_marker_leave_one_out.csv",
                    "results/bam_robustness/GSE286068_BAM_marker_leave_one_out.csv",
                ), "Marker leave-one-out BAM-like annotation audit."),
                SheetSpec("Resolution_sensitivity", (
                    "results/bam_robustness/GSE138266_BAM_multiresolution_overlap.csv",
                    "results/bam_robustness/GSE194078_BAM_multiresolution_overlap.csv",
                    "results/bam_robustness/GSE286068_BAM_multiresolution_overlap.csv",
                ), "BAM-like overlap across Leiden resolutions."),
                SheetSpec("Tissue_restriction", "results/bam_robustness/BAM_tissue_restriction.csv", "CSF-versus-blood restriction of the stable BAM-like core."),
                SheetSpec("State_abundance_samples", "results/state_abundance/CSF_lineage_state_abundance_sample_table.csv", "Sample-level CSF myeloid lineage-state abundance table."),
                SheetSpec("State_abundance_tests", "results/state_abundance/CSF_lineage_state_abundance_cohort_tests.csv", "Within-cohort disease tests of lineage-state abundance."),
                SheetSpec("State_abundance_meta", "results/state_abundance/CSF_lineage_state_abundance_meta_tests.csv", "Cross-cohort abundance meta-analysis."),
            ),
        ),
        WorkbookSpec(
            "Supplementary_Table_3_Trajectory_TwoAxis_LAM.xlsx",
            "Supplementary Table 3 | Trajectory, two-axis decoupling and cross-species LAM programs",
            "Trajectory robustness, donor-level differentiation/inflammation coordinates and ortholog-mapped LAM program tests.",
            (
                SheetSpec("Continuum_sampling", "results/trajectory/continuum_sampling_ledger.csv", "Cells and donors eligible for continuum inference."),
                SheetSpec("Continuum_states", "results/trajectory/continuum_state_summary.csv", "State-level trajectory summary."),
                SheetSpec("Trajectory_topology", "results/trajectory/continuum_topology_qc.csv", "Root-to-terminal topology audit."),
                SheetSpec("Method_concordance", "results/trajectory/consensus_method_concordance.csv", "Within-donor concordance of trajectory representations."),
                SheetSpec("Donor_jackknife", "results/trajectory/trajectory_donor_jackknife.csv", "Leave-one-donor trajectory reconstruction audit."),
                SheetSpec("Multiroot_audit", "results/trajectory/multiroot_DPT_qc.csv", "Sensitivity to a multi-donor root definition."),
                SheetSpec("Two_axis_donors", "results/two_axis/two_axis_donor_scores.csv", "Donor-level differentiation, activation residual and decoupling scores."),
                SheetSpec("Two_axis_disease", "results/two_axis/two_axis_disease_tests.csv", "Within-cohort disease tests for the two-axis endpoints."),
                SheetSpec("Two_axis_meta", "results/two_axis/two_axis_meta_tests.csv", "Cross-cohort two-axis meta-analysis."),
                SheetSpec("Threshold_grid", "results/two_axis/decoupling_threshold_meta_grid.csv", "Prespecified decoupling-threshold sensitivity grid."),
                SheetSpec("LAM_signature_ledger", "results/cross_species_lam/meningeal_LAM_signature_ledger.csv", "Mouse program membership and human ortholog ledger."),
                SheetSpec("LAM_availability", "results/cross_species_lam/LAM_signature_availability.csv", "Per-cohort availability of ortholog-mapped LAM genes."),
                SheetSpec("LAM_donor_scores", "results/cross_species_lam/LAM_donor_scores.csv", "Donor-level raw and differentiation-adjusted LAM program scores."),
                SheetSpec("LAM_pair_tests", "results/cross_species_lam/LAM_decoupled_pair_tests.csv", "Within-cohort donor-paired decoupled-state LAM tests."),
                SheetSpec("LAM_meta", "results/cross_species_lam/LAM_decoupled_meta.csv", "Cross-cohort LAM program meta-analysis."),
            ),
        ),
        WorkbookSpec(
            "Supplementary_Table_4_Receiver_Sender_Regulators.xlsx",
            "Supplementary Table 4 | Receptor, ligand and regulatory associations",
            "Receptor availability, canonical ligand-receptor circuits, target-held-out regulons and donor-level association tests.",
            (
                SheetSpec("Receptor_availability", "results/niche_regulator/receiver_receptor_availability.csv", "Per-cohort availability of prespecified receiver complexes."),
                SheetSpec("Canonical_circuits", "results/niche_regulator/LIANA_consensus_canonical_circuit_ledger.csv", "Canonical ligand-receptor circuit evidence ledger."),
                SheetSpec("Receiver_ranking", "results/niche_regulator/receiver_side_circuit_ranking.csv", "Integrated ranking of receiver-side circuits."),
                SheetSpec("Receiver_delta_meta", "results/niche_regulator/receiver_state_delta_meta.csv", "Cross-cohort receiver-state delta meta-analysis."),
                SheetSpec("Regulator_ranking", "results/niche_regulator/directional_regulator_candidate_ranking.csv", "Direction-aware regulator candidate ranking."),
                SheetSpec("Regulator_holdout", "results/niche_regulator/regulator_noncircularity_audit.csv", "Target-removal audit preventing endpoint-regulon circularity."),
                SheetSpec("Regulator_LAM_assoc", "results/niche_regulator/regulator_delta_to_LAM_outcome_associations.csv", "Donor-level regulator delta to residualized LAM outcome associations."),
                SheetSpec("Sender_gene_measure", "results/niche_regulator/CSF_T_sender_gene_measurability.csv", "Per-cohort measurability of prespecified CSF T-cell sender genes."),
                SheetSpec("Sender_donor_scores", "results/niche_regulator/CSF_T_sender_donor_scores.csv", "Donor-level CSF T-cell sender-family scores."),
                SheetSpec("Sender_LAM_assoc", "results/niche_regulator/sender_to_LAM_outcome_associations.csv", "Donor-level sender-family to residualized LAM outcome tests."),
            ),
        ),
        WorkbookSpec(
            "Supplementary_Table_5_InVivo_CRISPR_Bridge.xlsx",
            "Supplementary Table 5 | In vivo CRISPR signature bridge",
            "External mouse perturbation provenance, ortholog mapping and donor-level human directional tests.",
            (
                SheetSpec("External_sources", "results/in_vivo_crispr_bridge/external_source_manifest.csv", "External experimental source files and checksums."),
                SheetSpec("Source_QC", "results/in_vivo_crispr_bridge/external_signature_source_qc.csv", "Direct-KO cross-model concordance and source-cleaning audit."),
                SheetSpec("Direct_KO_ledger", "results/in_vivo_crispr_bridge/direct_ko_consensus_signature_ledger.csv", "Balanced direct-KO consensus signature ledger."),
                SheetSpec("Response_sig_ledger", "results/in_vivo_crispr_bridge/published_cytokine_signature_ledger.csv", "Published cytokine-response signature ledger."),
                SheetSpec("Ortholog_map", "results/in_vivo_crispr_bridge/mouse_human_one_to_one_orthologs.csv", "Mouse-human one-to-one ortholog map used by the bridge."),
                SheetSpec("Human_availability", "results/in_vivo_crispr_bridge/human_signature_availability.csv", "Per-cohort mapped signature availability."),
                SheetSpec("Residualization_QC", "results/in_vivo_crispr_bridge/human_score_residualization_diagnostics.csv", "Control-manifold residualization diagnostics."),
                SheetSpec("Donor_effects", "results/in_vivo_crispr_bridge/donor_level_bridge_effects.csv", "Donor-level response and direct-KO-like bridge effects."),
                SheetSpec("Cohort_tests", "results/in_vivo_crispr_bridge/cohort_bridge_tests.csv", "Within-cohort directional sign-flip tests."),
                SheetSpec("Bridge_meta", "results/in_vivo_crispr_bridge/bridge_meta_analysis.csv", "Cross-cohort Stouffer meta-analysis."),
                SheetSpec("Primary_decisions", "results/in_vivo_crispr_bridge/primary_bridge_decisions.csv", "Prespecified perturbation-level advancement decisions."),
                SheetSpec("Leave_one_cohort_out", "results/in_vivo_crispr_bridge/leave_one_cohort_out_audit.csv", "Leave-one-cohort-out bridge robustness audit."),
            ),
        ),
        WorkbookSpec(
            "Supplementary_Table_6_Virtual_Perturbation.xlsx",
            "Supplementary Table 6 | Virtual perturbation specificity, veto and rescue tests",
            "Observed-GRN directional tests, target-shuffled network controls, endpoint holdouts and combination rescue audits.",
            (
                SheetSpec("Endpoint_holdouts", "results/virtual_perturbation_celloracle/all_endpoint_holdout_ledger.csv", "Genes removed from candidate GRNs to prevent endpoint leakage."),
                SheetSpec("Perturbation_status", "results/virtual_perturbation_celloracle/all_perturbation_status.csv", "Per-cohort candidate perturbation execution status."),
                SheetSpec("Primary_cohort", "results/virtual_perturbation_celloracle/primary_endpoint_cohort_ledger.csv", "Within-cohort primary virtual perturbation endpoints."),
                SheetSpec("Primary_meta", "results/virtual_perturbation_celloracle/primary_endpoint_meta_analysis.csv", "Cross-cohort expected-direction and network-specificity meta-analysis."),
                SheetSpec("Donor_signflip", "results/virtual_perturbation_celloracle/observed_GRN_donor_signflip_tests.csv", "Donor-level observed-GRN sign-flip tests."),
                SheetSpec("Shuffled_networks", "results/virtual_perturbation_celloracle/shuffled_GRN_empirical_tests.csv", "Target-shuffled GRN empirical null distributions."),
                SheetSpec("Harm_veto", "results/virtual_perturbation_celloracle/harm_veto_meta_analysis.csv", "Complete prespecified harmful-direction family."),
                SheetSpec("Directional_decisions", "results/virtual_perturbation_celloracle/directional_candidate_decisions.csv", "Candidate advancement decisions after direction and veto checks."),
                SheetSpec("Combination_status", "results/virtual_perturbation_celloracle_combinations/all_combination_status.csv", "Per-cohort combination execution status."),
                SheetSpec("Combination_cohort", "results/virtual_perturbation_celloracle_combinations/combination_cohort_tests.csv", "Within-cohort combination directional tests."),
                SheetSpec("Combination_meta", "results/virtual_perturbation_celloracle_combinations/combination_criterion_meta_analysis.csv", "Cross-cohort combination rescue meta-analysis."),
                SheetSpec("Combination_decisions", "results/virtual_perturbation_celloracle_combinations/combination_decisions.csv", "Combination advancement decisions."),
                SheetSpec("Combination_shuffle", "results/virtual_perturbation_celloracle_combinations/combination_shuffled_GRN_rescue_tests.csv", "Shuffled-network combination rescue controls."),
            ),
        ),
        WorkbookSpec(
            "Supplementary_Table_7_Drug_Target_Structure.xlsx",
            "Supplementary Table 7 | Transition-aware drug, target and structure audit",
            "Frozen phenotype signatures, complete chemical screen, genetic target deconvolution, mechanism neighbourhood and docking validation.",
            (
                SheetSpec("Signature_build", "results/transition_aware_drug_screen/signatures/signature_build_summary.csv", "Frozen transition-aware signature construction summary."),
                SheetSpec("Inflammatory_activation", "results/transition_aware_drug_screen/signatures/inflammatory_activation_signature.csv", "Fixed inflammatory activation gene weights."),
                SheetSpec("Donor_eligibility", "results/transition_aware_drug_screen/signatures/donor_pair_eligibility.csv", "Donor-state pairs eligible for signature construction."),
                SheetSpec("NIBR_profile_screen", "results/transition_aware_drug_screen/nibr_full_profiles/NIBR_transition_candidate_profile_scores.csv", "Complete 76-profile NIBR candidate screen."),
                SheetSpec("NIBR_compounds", "results/transition_aware_drug_screen/nibr_full_profiles/NIBR_transition_candidate_compound_summary.csv", "Compound-level NIBR profile summary."),
                SheetSpec("Dose_consistency", "results/transition_aware_drug_screen/nibr_full_profiles/NIBR_transition_candidate_dose_consistency.csv", "Cross-dose direction and advancement audit."),
                SheetSpec("Screen_summary", "results/transition_aware_drug_screen/directional_screen/directional_screen_summary.csv", "Cross-library directional screen summary."),
                SheetSpec("All_compounds", "results/transition_aware_drug_screen/directional_screen/all_compound_scores.csv", "All compound-level directional scores across chemical libraries."),
                SheetSpec("Genetic_coverage", "results/transition_aware_drug_screen/genetic_deconvolution/genetic_target_coverage.csv", "Target representation across loss-of-function libraries."),
                SheetSpec("Genetic_library", "results/transition_aware_drug_screen/genetic_deconvolution/genetic_target_library_summary.csv", "Per-target, per-library genetic directional summary."),
                SheetSpec("Genetic_decisions", "results/transition_aware_drug_screen/genetic_deconvolution/genetic_target_cross_library_decisions.csv", "Cross-library genetic target decisions."),
                SheetSpec("BX912_neighbors", "results/transition_aware_drug_screen/mechanism_deconvolution/BX912_compound_level_directional_neighbors.csv", "Unbiased compound-level directional neighbourhood of BX-912."),
                SheetSpec("Neighbor_enrichment", "results/transition_aware_drug_screen/mechanism_deconvolution/BX912_top5pct_neighbor_annotation_enrichment.csv", "Target/MOA enrichment among the top BX-912 neighbours."),
                SheetSpec("Redocking_QC", "results/docking/1Z5M_PDPK1_BX320_site/1Z5M_BX320_redocking_validation.csv", "Symmetry-aware control-ligand redocking validation."),
                SheetSpec("Docking_grid", "results/docking/1Z5M_PDPK1_BX320_site/1Z5M_BX320_grid_definition.csv", "Validated PDPK1 docking-grid definition."),
                SheetSpec("Receptor_prep", "results/docking/1Z5M_PDPK1_BX320_site/1Z5M_receptor_preparation_decisions.csv", "PDPK1 receptor preparation decisions."),
                SheetSpec("BX912_pose_summary", "results/docking/1Z5M_PDPK1_BX912/BX912_PDPK1_pose_summary.csv", "BX-912 docking pose summary."),
                SheetSpec("BX912_robustness", "results/docking/1Z5M_PDPK1_BX912/BX912_pose_robustness.csv", "BX-912 pose recovery across seeds and search modes."),
                SheetSpec("BX912_contacts", "results/docking/1Z5M_PDPK1_BX912/BX912_PDPK1_contacts.csv", "Residue-level BX-912 pose contacts."),
            ),
        ),
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=ROOT)
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_sheet(root: Path, spec: SheetSpec) -> tuple[pd.DataFrame, str]:
    if isinstance(spec.source, tuple):
        frame = combine_with_dataset(root, spec.source)
        source_text = "; ".join(spec.source)
    else:
        frame = pd.read_csv(root / spec.source)
        source_text = spec.source
    if spec.transform:
        frame = spec.transform(frame)
    return frame, source_text


def excel_value(value: object) -> object:
    if pd.isna(value):
        return None
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    if hasattr(value, "item"):
        return value.item()
    return value


def number_format(column: str, series: pd.Series) -> str:
    lower = column.lower()
    if not pd.api.types.is_numeric_dtype(series):
        return "General"
    if pd.api.types.is_integer_dtype(series.dropna()):
        return "0"
    if any(token in lower for token in ("p_value", "pvalue", "_p", "fdr", "q_value")):
        return "0.00E+00"
    return "0.000"


def format_overview(ws, spec: WorkbookSpec, sheet_rows: list[dict[str, object]]) -> None:
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = spec.title
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A3:F4")
    ws["A3"] = spec.scope
    ws["A3"].font = Font(name="Arial", size=10, color="26333D")
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A6"] = "Workbook rules"
    ws["A6"].font = Font(name="Arial", size=11, bold=True, color=NAVY)
    rules = (
        "Each data sheet is a static export from the listed source file; no values are recalculated in Excel. "
        "Blank cells represent unavailable or non-applicable values. Filters and frozen headers are enabled."
    )
    ws.merge_cells("A7:F8")
    ws["A7"] = rules
    ws["A7"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A7"].font = Font(name="Arial", size=9, color=MID_GREY)
    headers = ["Sheet", "Description", "Source file(s)", "Rows", "Columns", "Source SHA-256"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(10, col, value)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row_idx, item in enumerate(sheet_rows, 11):
        values = [item[key] for key in ("sheet", "description", "source", "rows", "columns", "source_sha256")]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_idx, col, value)
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(wrap_text=col in {2, 3, 6}, vertical="top")
            if row_idx % 2:
                cell.fill = PatternFill("solid", fgColor=PALE_GREY)
        ws.cell(row_idx, 1).hyperlink = f"#'{item['sheet']}'!A1"
        ws.cell(row_idx, 1).style = "Hyperlink"
    ws.freeze_panes = "A11"
    ws.auto_filter.ref = f"A10:F{10 + len(sheet_rows)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_area = f"A1:F{10 + len(sheet_rows)}"
    widths = [26, 62, 70, 10, 10, 66]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def format_data_sheet(ws, spec: SheetSpec, frame: pd.DataFrame, source_text: str) -> None:
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(frame.columns)))
    ws.cell(1, 1, spec.description)
    ws.cell(1, 1).font = Font(name="Arial", size=11, bold=True, color=WHITE)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(1, len(frame.columns)))
    ws.cell(2, 1, f"Source: {source_text}")
    ws.cell(2, 1).font = Font(name="Arial", size=8, italic=True, color=MID_GREY)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 26
    thin = Side(style="thin", color="D6DCE0")
    for col_idx, column in enumerate(frame.columns, 1):
        cell = ws.cell(4, col_idx, str(column))
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=thin)
    ws.row_dimensions[4].height = 30
    formats = {column: number_format(str(column), frame[column]) for column in frame.columns}
    for row_idx, values in enumerate(frame.itertuples(index=False, name=None), 5):
        for col_idx, (column, value) in enumerate(zip(frame.columns, values), 1):
            cell = ws.cell(row_idx, col_idx, excel_value(value))
            cell.font = Font(name="Arial", size=8)
            cell.alignment = Alignment(vertical="top")
            cell.number_format = formats[column]
    max_row = 4 + len(frame)
    max_col = len(frame.columns)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(max_col)}{max_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"
    ws.print_area = f"A1:{get_column_letter(max_col)}{max_row}"
    table_name = "T_" + "".join(ch if ch.isalnum() else "_" for ch in ws.title)
    table = Table(displayName=table_name[:250], ref=ws.auto_filter.ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)
    for col_idx, column in enumerate(frame.columns, 1):
        sample = [str(column)] + [str(v) for v in frame[column].dropna().head(250)]
        width = min(42, max(9, max(len(value) for value in sample) + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def create_workbook(root: Path, output: Path, spec: WorkbookSpec) -> list[dict[str, object]]:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    sheet_rows = []
    for sheet_spec in spec.sheets:
        frame, source_text = read_sheet(root, sheet_spec)
        ws = workbook.create_sheet(sheet_spec.name)
        format_data_sheet(ws, sheet_spec, frame, source_text)
        source_hashes = []
        source_paths = sheet_spec.source if isinstance(sheet_spec.source, tuple) else (sheet_spec.source,)
        for source_path in source_paths:
            source_hashes.append(sha256(root / source_path))
        sheet_rows.append({
            "sheet": sheet_spec.name,
            "description": sheet_spec.description,
            "source": source_text,
            "rows": len(frame),
            "columns": len(frame.columns),
            "source_sha256": "; ".join(source_hashes),
        })
    format_overview(overview, spec, sheet_rows)
    path = output / spec.filename
    workbook.save(path)
    return [dict(item, workbook=spec.filename) for item in sheet_rows]


def write_index(output: Path, workbook_specs: tuple[WorkbookSpec, ...], manifest: pd.DataFrame) -> None:
    lines = [
        "# Supplementary tables index", "", "Generated: 2026-09-03", "",
        "Each workbook contains an Overview sheet with scope, source paths, row counts and source-file SHA-256 digests. Data sheets are static, filterable exports with frozen headers; no statistical value is recalculated in Excel.", "",
    ]
    for number, spec in enumerate(workbook_specs, 1):
        rows = manifest.loc[manifest.workbook.eq(spec.filename)]
        lines.extend([
            f"## Supplementary Table {number}", "",
            f"**{spec.title.split('|', 1)[-1].strip()}**", "", spec.scope, "",
            f"Workbook: `{spec.filename}`", "",
            f"Sheets: {', '.join(rows.sheet)}.", "",
        ])
    (output / "SUPPLEMENTARY_TABLES_INDEX.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    args = parse_args()
    root = args.root.resolve()
    output = root / "supplementary"
    output.mkdir(parents=True, exist_ok=True)
    workbook_specs = specs()
    records = []
    for workbook_spec in workbook_specs:
        records.extend(create_workbook(root, output, workbook_spec))
        print(f"Built {workbook_spec.filename}", flush=True)
    source_manifest = pd.DataFrame(records)
    source_manifest.to_csv(output / "SUPPLEMENTARY_SOURCE_MANIFEST.csv", index=False)
    write_index(output, workbook_specs, source_manifest)
    workbook_manifest = []
    for workbook_spec in workbook_specs:
        path = output / workbook_spec.filename
        workbook_manifest.append({
            "workbook": workbook_spec.filename,
            "sheets": len(workbook_spec.sheets) + 1,
            "size_bytes": path.stat().st_size,
            "sha256": sha256(path),
            "generated_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        })
    pd.DataFrame(workbook_manifest).to_csv(output / "SUPPLEMENTARY_WORKBOOK_MANIFEST.csv", index=False)


if __name__ == "__main__":
    main()
