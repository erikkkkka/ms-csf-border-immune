#!/usr/bin/env python3
"""Audit supplementary workbooks, source hashes and spreadsheet usability."""

from __future__ import annotations

import argparse
import hashlib
import importlib
import math
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
EXPECTED = {
    "Supplementary_Table_1_Cohorts_and_QC.xlsx": 11,
    "Supplementary_Table_2_Annotation_and_State.xlsx": 11,
    "Supplementary_Table_3_Trajectory_TwoAxis_LAM.xlsx": 19,
    "Supplementary_Table_4_Receiver_Sender_Regulators.xlsx": 11,
    "Supplementary_Table_5_InVivo_CRISPR_Bridge.xlsx": 13,
    "Supplementary_Table_6_Virtual_Perturbation.xlsx": 14,
    "Supplementary_Table_7_Drug_Target_Structure.xlsx": 20,
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=ROOT)
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def audit_workbooks(root: Path) -> pd.DataFrame:
    output = root / "supplementary"
    rows: list[dict[str, object]] = []
    for filename, expected_sheets in EXPECTED.items():
        path = output / filename
        row: dict[str, object] = {
            "workbook": filename,
            "exists": path.is_file(),
            "expected_sheets": expected_sheets,
        }
        if not path.is_file():
            row["passed"] = False
            rows.append(row)
            continue
        workbook = load_workbook(path, read_only=False, data_only=False)
        formula_cells = 0
        error_cells = 0
        data_sheets_ok = True
        table_count = 0
        arial_headers = True
        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            for row_cells in ws.iter_rows():
                for cell in row_cells:
                    formula_cells += int(cell.data_type == "f")
                    error_cells += int(cell.data_type == "e")
            if sheet_name == "Overview":
                data_sheets_ok &= ws.freeze_panes == "A11" and bool(ws.auto_filter.ref)
                arial_headers &= ws["A1"].font.name == "Arial"
            else:
                data_sheets_ok &= (
                    ws.max_row >= 5
                    and ws.max_column >= 1
                    and ws.freeze_panes == "A5"
                    and bool(ws.auto_filter.ref)
                    and len(ws.tables) == 1
                )
                table_count += len(ws.tables)
                arial_headers &= all(ws.cell(4, col).font.name == "Arial" for col in range(1, ws.max_column + 1))
        row.update({
            "actual_sheets": len(workbook.sheetnames),
            "data_sheets": len(workbook.sheetnames) - 1,
            "table_count": table_count,
            "formula_cells": formula_cells,
            "error_cells": error_cells,
            "freeze_filter_table_checks": bool(data_sheets_ok),
            "arial_header_checks": bool(arial_headers),
            "size_bytes": path.stat().st_size,
            "sha256": sha256(path),
        })
        row["passed"] = bool(
            len(workbook.sheetnames) == expected_sheets
            and workbook.sheetnames[0] == "Overview"
            and data_sheets_ok
            and arial_headers
            and formula_cells == 0
            and error_cells == 0
            and path.stat().st_size > 10_000
        )
        rows.append(row)
    return pd.DataFrame(rows)


def audit_sources(root: Path) -> pd.DataFrame:
    manifest_path = root / "supplementary/SUPPLEMENTARY_SOURCE_MANIFEST.csv"
    manifest = pd.read_csv(manifest_path)
    rows: list[dict[str, object]] = []
    for entry in manifest.itertuples(index=False):
        source_paths = [part.strip() for part in str(entry.source).split(";")]
        expected_hashes = [part.strip() for part in str(entry.source_sha256).split(";")]
        actual_hashes = [sha256(root / source_path) if (root / source_path).is_file() else "MISSING" for source_path in source_paths]
        rows.append({
            "workbook": entry.workbook,
            "sheet": entry.sheet,
            "source_files": len(source_paths),
            "reported_rows": int(entry.rows),
            "reported_columns": int(entry.columns),
            "sources_exist": all(value != "MISSING" for value in actual_hashes),
            "source_hashes_match": actual_hashes == expected_hashes,
            "passed": all(value != "MISSING" for value in actual_hashes) and actual_hashes == expected_hashes,
        })
    return pd.DataFrame(rows)


def read_result_sheet(path: Path, sheet: str) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        worksheet = workbook[sheet]
        headers = [cell.value for cell in worksheet[4]]
        return pd.DataFrame(worksheet.iter_rows(min_row=5, values_only=True), columns=headers)
    finally:
        workbook.close()


def matches_result_source(path: Path, sheet: str, source: pd.DataFrame) -> bool:
    try:
        actual = read_result_sheet(path, sheet)
        pd.testing.assert_frame_equal(actual, source, check_dtype=False, check_exact=False, rtol=1e-12, atol=1e-14)
    except (AssertionError, KeyError):
        return False
    return True


def audit_result_consistency(root: Path) -> pd.DataFrame:
    table3 = root / "supplementary/Supplementary_Table_3_Trajectory_TwoAxis_LAM.xlsx"
    primary = pd.read_csv(root / "results/external_validation/GSE328265_two_axis_projection/three_cohort_two_axis_meta.csv")
    discovery = pd.read_csv(root / "results/two_axis/two_axis_meta_tests.csv")
    endpoints = {"continuous_decoupling_index", "decoupled_inflammatory_fraction"}
    rows = [
        {
            "check": "primary_three_cohort_meta",
            "detail": "Two_axis_meta must match the three-cohort, two-endpoint synthesis.",
            "passed": bool(set(primary.endpoint) == endpoints and primary.n_cohorts.eq(3).all() and matches_result_source(table3, "Two_axis_meta", primary)),
        },
        {
            "check": "discovery_only_meta_retained",
            "detail": "Two_axis_meta_discovery retains the separate two-cohort, 11-test family.",
            "passed": bool(len(discovery) == 11 and discovery.n_cohorts.eq(2).all() and matches_result_source(table3, "Two_axis_meta_discovery", discovery)),
        },
    ]
    full_auc = pd.read_csv(root / "results/trajectory/full_multiroot_endpoint_auc.csv")
    saved_auc = importlib.import_module("84_export_multiroot_endpoint_auc").full_multiroot_endpoints(root)
    try:
        pd.testing.assert_frame_equal(full_auc, saved_auc, check_dtype=False, check_exact=False, rtol=1e-12, atol=1e-14)
        saved_auc_matches = True
    except AssertionError:
        saved_auc_matches = False
    primary_auc = full_auc.loc[full_auc.representation.eq("primary")].set_index("dataset")
    plotted = pd.read_csv(root / "figures_main/source_data/Fig1e_trajectory_robustness.csv")
    jackknife = pd.read_csv(root / "results/trajectory/trajectory_donor_jackknife.csv")
    rows.append({
        "check": "full_multiroot_auc_saved_cells",
        "detail": "Six full-cohort AUCs match saved donor-root DPT scores and Table 3.",
        "passed": bool(len(full_auc) == 6 and saved_auc_matches
                       and matches_result_source(table3, "Multiroot_endpoint_AUC", full_auc)),
    })
    for entry in plotted.itertuples():
        subset = jackknife.loc[jackknife.dataset.eq(entry.dataset) & jackknife.representation.eq("primary")]
        values = subset.rebuilt_terminal_vs_root_auc
        rows.append({
            "check": f"Figure1e_same_method_{entry.dataset}",
            "detail": "Full and donor-omission AUCs use primary multiroot DPT and the same trajectory-role endpoints.",
            "passed": bool(entry.method == "multiroot DPT" and entry.representation == "primary"
                           and entry.full_score_column == "dpt_multiroot_primary"
                           and math.isclose(entry.full_auc, primary_auc.loc[entry.dataset, "full_terminal_vs_root_auc"], abs_tol=1e-12)
                           and entry.n_leave_one_donor_runs == len(values)
                           and all(math.isclose(a, b, abs_tol=1e-12) for a, b in (
                               (entry.jackknife_min, values.min()), (entry.jackknife_median, values.median()),
                               (entry.jackknife_max, values.max())))),
        })
    external = pd.read_csv(root / "results/external_validation/GSE328265_two_axis_projection/external_projection_topology.csv")
    rows.append({
        "check": "external_projection_state_order_and_endpoint_auc",
        "detail": "External validation fields are state_order_spearman and terminal_vs_root_auc, exported in Table 3.",
        "passed": bool(matches_result_source(table3, "External_projection_topology", external)
                       and len(external) == 1 and int(external.iloc[0].n_cells) == 3333
                       and math.isclose(external.iloc[0].state_order_spearman, 0.5949148474253424, abs_tol=1e-12)
                       and math.isclose(external.iloc[0].terminal_vs_root_auc, 0.9982686731146774, abs_tol=1e-12)),
    })
    table1 = root / "supplementary/Supplementary_Table_1_Cohorts_and_QC.xlsx"
    roles = pd.read_csv(root / "metadata/cohort_analysis_roles.csv")
    rows.append({
        "check": "cohort_analysis_roles",
        "detail": "Two disease discovery cohorts, three trajectory cohorts and one external projection cohort.",
        "passed": bool(matches_result_source(table1, "Cohort_analysis_roles", roles)
                       and set(roles.loc[roles.disease_discovery, "dataset"]) == {"GSE138266", "GSE286068"}
                       and set(roles.loc[roles.independent_trajectory_reconstruction, "dataset"]) == {"GSE138266", "GSE194078", "GSE286068"}
                       and set(roles.loc[roles.external_fixed_score_projection, "dataset"]) == {"GSE328265"}),
    })
    postdecont = read_result_sheet(table1, "Postdecont_expression_QC")
    for dataset, expected in (("GSE194078", (137903, 137108, 795)), ("GSE286068", (22825, 22756, 69))):
        source = pd.read_csv(root / f"results/qc/{dataset}_postdecont_qc_summary.csv")
        actual = postdecont.loc[postdecont.dataset.eq(dataset)].drop(columns="dataset").reset_index(drop=True)
        totals = tuple(int(source[column].sum()) for column in (
            "n_postdoublet_input", "n_postdecont_qc_pass", "n_postdecont_removed"))
        try:
            pd.testing.assert_frame_equal(actual, source, check_dtype=False)
            matched = True
        except AssertionError:
            matched = False
        rows.append({
            "check": f"postdecont_expression_filter_{dataset}",
            "detail": f"Input / retained / removed counts: {totals}; separate from doublet-union singlets.",
            "passed": bool(matched and totals == expected and totals[0] - totals[1] == totals[2]),
        })
    effect = pd.read_csv(root / "results/cross_species_lam/LAM_decoupled_meta.csv")
    precursor = effect.loc[effect.endpoint.eq("decoupled_excess__mouse_GMCSF_LAM_precursor")].iloc[0]
    rows.append({
        "check": "precursor_three_decimal_rounding",
        "detail": "Mean cohort precursor excess rounds directly to 0.201 at three decimals.",
        "passed": f"{precursor.mean_cohort_excess:.3f}" == "0.201",
    })
    grn = pd.read_csv(root / "figures_extended/source_data/EDFig8a_observed_vs_network_specific.csv")
    irf1 = grn.loc[grn.candidate.eq("IRF1") & grn.endpoint.eq("axis2")].iloc[0]
    rows.append({
        "check": "cohort_GRN_and_distinct_IRF1_tests",
        "detail": "Whole-cohort GRNs; expected-direction and shuffled-network specificity statistics remain distinct.",
        "passed": bool(grn.evidence_scope.eq("virtual perturbation in fitted whole-cohort GRNs; effects aggregated by donor").all()
                       and math.isclose(irf1.combined_expected_z, 5.685229521351445, abs_tol=1e-12)
                       and math.isclose(irf1.combined_network_specific_z, 0.8429068483277294, abs_tol=1e-12)),
    })
    table5 = root / "supplementary/Supplementary_Table_5_InVivo_CRISPR_Bridge.xlsx"
    qc = read_result_sheet(table5, "Source_QC").set_index("perturbation")
    published = pd.read_csv(root / "results/in_vivo_crispr_bridge/published_cytokine_signature_ledger.csv")
    for perturbation, frame in published.groupby("perturbation", observed=True):
        mapped = frame.human_gene.fillna("").ne("")
        negative = frame.hoxb8_log2fc.lt(0) & frame.chimera_log2fc.lt(0)
        entry = qc.loc[perturbation]
        fraction = float(negative.mean())
        rows.append({
            "check": f"published_response_denominator_{perturbation}",
            "detail": f"{int(negative.sum())}/{len(frame)} published mouse genes; {int(mapped.sum())} one-to-one mapped genes are not the denominator.",
            "passed": bool(
                int(entry.n_published_response_genes) == len(frame)
                and int(entry.n_published_one_to_one_mapped) == int(mapped.sum())
                and math.isclose(float(entry.published_fraction_negative_in_both_KO_models), fraction, rel_tol=1e-12, abs_tol=1e-14)
            ),
        })
    return pd.DataFrame(rows)


def main() -> None:
    args = parse_args()
    root = args.root.resolve()
    output = root / "supplementary"
    workbook_qa = audit_workbooks(root)
    source_qa = audit_sources(root)
    consistency_qa = audit_result_consistency(root)
    workbook_qa.to_csv(output / "SUPPLEMENTARY_WORKBOOK_QA.csv", index=False)
    source_qa.to_csv(output / "SUPPLEMENTARY_SOURCE_QA.csv", index=False)
    consistency_qa.to_csv(output / "SUPPLEMENTARY_CONSISTENCY_QA.csv", index=False)

    recorded = pd.read_csv(output / "SUPPLEMENTARY_WORKBOOK_MANIFEST.csv")
    recorded_hashes = dict(zip(recorded.workbook, recorded.sha256))
    manifest_match = all(recorded_hashes.get(row.workbook) == row.sha256 for row in workbook_qa.itertuples())
    index_exists = (output / "SUPPLEMENTARY_TABLES_INDEX.md").is_file()
    overall_pass = bool(
        workbook_qa.passed.all() and source_qa.passed.all() and consistency_qa.passed.all() and manifest_match and index_exists
    )
    total_data_sheets = int(workbook_qa.data_sheets.sum())
    report = f"""# Supplementary table QA report

Audited at: {datetime.now(timezone.utc).isoformat(timespec='seconds')}

## Status

**{'PASS' if overall_pass else 'FAIL'}**

- Workbooks: {int(workbook_qa.passed.sum())}/{len(workbook_qa)} passed.
- Data sheets: {total_data_sheets} across seven workbooks.
- Source links and hashes: {int(source_qa.passed.sum())}/{len(source_qa)} passed.
- Result scope, figure baselines and response-gene denominators: {int(consistency_qa.passed.sum())}/{len(consistency_qa)} passed.
- Workbook manifest hashes: {'matched' if manifest_match else 'mismatch'}.
- Formula cells: {int(workbook_qa.formula_cells.sum())}; Excel error cells: {int(workbook_qa.error_cells.sum())}.

## Checks

- Every workbook opens successfully and begins with an Overview sheet.
- Every data sheet is non-empty, uses Arial headers, and has a frozen header,
  autofilter and formatted Excel table.
- Data are static result exports; no spreadsheet formulas are introduced.
- Every declared source path exists and matches its recorded SHA-256 digest.
- Workbook digests match `SUPPLEMENTARY_WORKBOOK_MANIFEST.csv`.
- Primary three-cohort disease results and discovery-only results are exported in separate sheets.
- Response-gene percentages use all originally published mouse genes, before ortholog filtering.
- Figure 1e compares primary multiroot DPT with method-matched donor omissions; complete AUCs are checked against saved cells.
- External projection topology, final cohort roles and post-decontamination expression-filter counts are exported separately.
- Cohort-GRN directional tests remain separate from shuffled-network specificity tests.
"""
    (output / "SUPPLEMENTARY_QA_REPORT.md").write_text(report, encoding="utf-8")
    print(report)
    if not overall_pass:
        raise SystemExit("Supplementary table QA failed")


if __name__ == "__main__":
    main()
